import html
import re
from functools import lru_cache
from io import BytesIO
from pathlib import Path
from typing import Dict, Iterable, List, Optional

from backend.app.schemas.preferences import ResumePreferences
from backend.app.services.preferences import paper_size
from backend.app.services.template_catalog import heading_background, template_style
from backend.app.schemas.candidate import CandidateProfile
from backend.app.schemas.strategy import ResumeDraft, ResumeSectionItem


DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PDF_MEDIA_TYPE = "application/pdf"


@lru_cache(maxsize=4)
def _flow_pdf_font(serif=False):
    """Embed an available Chinese font; retain the portable CID fallback."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont, TTFError
    path = Path("/System/Library/Fonts/Supplemental/Songti.ttc" if serif
                else "/System/Library/Fonts/STHeiti Light.ttc")
    name = "ResumeSongti" if serif else "ResumeHeiti"
    if path.is_file():
        try:
            pdfmetrics.registerFont(TTFont(name, str(path), subfontIndex=0 if serif else 1))
            return name
        except (OSError, TTFError):
            pass
    return "STSong-Light"


def sanitize_filename(value: str, fallback: str = "tailored-resume") -> str:
    name = (value or "").strip()
    name = re.sub(r"\.(?:docx|pdf)$", "", name, flags=re.IGNORECASE)
    name = re.sub(r'[\\/:*?"<>|\x00-\x1f]', "-", name)
    name = re.sub(r"\s+", " ", name).strip(" .-_")
    return (name[:80].rstrip(" .-_") or fallback)


def draft_to_plain_text(draft: ResumeDraft, candidate: CandidateProfile) -> str:
    labels = _labels(draft)
    lines: List[str] = []
    title = draft.headline.strip() or (candidate.basics.name or "").strip()
    if title:
        lines.append(title)
    contact = _contact_line(candidate)
    if contact:
        lines.append(contact)
    if draft.summary.strip():
        lines.extend(["", labels["summary"], draft.summary.strip()])
    if draft.skills_section:
        lines.extend(["", labels["skills"], ", ".join(_clean_items(draft.skills_section))])
    _append_plain_sections(lines, labels["experience"], draft.experience_section)
    _append_plain_sections(lines, labels["projects"], draft.project_section)
    education = _clean_items(draft.education_section)
    if education:
        lines.extend(["", labels["education"]])
        lines.extend("- {0}".format(item) for item in education)
    return "\n".join(lines).strip()


def build_docx(draft: ResumeDraft, candidate: CandidateProfile, preferences: Optional[ResumePreferences] = None) -> bytes:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Mm, Pt, RGBColor

    preferences = preferences or ResumePreferences()
    compact = preferences.template == "ats_compact"
    layout = template_style(preferences.template)
    document = Document()
    section = document.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    if paper_size(preferences) == "A4":
        section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)

    styles = document.styles
    normal = styles["Normal"]
    docx_font_name = "Heiti SC"
    normal.font.name = docx_font_name
    normal.font.size = Pt(10 if compact else 10.25)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal._element.rPr.rFonts.set(qn("w:ascii"), docx_font_name)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), docx_font_name)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), docx_font_name)
    normal.paragraph_format.space_after = Pt(1 if compact else 2)
    normal.paragraph_format.line_spacing = 1.0

    title_style = styles["Title"]
    title_style.font.name = docx_font_name
    title_style.font.size = Pt(16)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style._element.rPr.rFonts.set(qn("w:ascii"), docx_font_name)
    title_style._element.rPr.rFonts.set(qn("w:hAnsi"), docx_font_name)
    title_style._element.rPr.rFonts.set(qn("w:eastAsia"), docx_font_name)
    title_style.paragraph_format.space_after = Pt(2)
    title_properties = title_style._element.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)
    title_run_properties = title_style._element.get_or_add_rPr()
    title_character_spacing = title_run_properties.find(qn("w:spacing"))
    if title_character_spacing is not None:
        title_run_properties.remove(title_character_spacing)

    if "Resume Section" not in styles:
        section_style = styles.add_style("Resume Section", WD_STYLE_TYPE.PARAGRAPH)
    else:
        section_style = styles["Resume Section"]
    section_style.font.name = docx_font_name
    section_style.font.size = Pt(11.5)
    section_style.font.bold = True
    section_style.font.color.rgb = RGBColor(0, 0, 0)
    section_style._element.rPr.rFonts.set(qn("w:ascii"), docx_font_name)
    section_style._element.rPr.rFonts.set(qn("w:hAnsi"), docx_font_name)
    section_style._element.rPr.rFonts.set(qn("w:eastAsia"), docx_font_name)
    section_style.paragraph_format.space_before = Pt(4 if compact else 6)
    section_style.paragraph_format.space_after = Pt(2)
    section_style.paragraph_format.keep_with_next = True

    bullet_style = styles["List Bullet"]
    bullet_style.font.name = docx_font_name
    bullet_style.font.size = Pt(10)
    bullet_style.font.color.rgb = RGBColor(0, 0, 0)
    bullet_style._element.rPr.rFonts.set(qn("w:ascii"), docx_font_name)
    bullet_style._element.rPr.rFonts.set(qn("w:hAnsi"), docx_font_name)
    bullet_style._element.rPr.rFonts.set(qn("w:eastAsia"), docx_font_name)

    if layout["adaptation"] == "flow":
        section.left_margin = section.right_margin = Pt(layout["margin_x"])
        section.top_margin = section.bottom_margin = Pt(layout["margin_y"])
        for style in (normal, title_style, section_style, bullet_style):
            style.font.name = layout["font"]
            for script in ("ascii", "hAnsi", "eastAsia"):
                style._element.get_or_add_rPr().rFonts.set(qn("w:" + script), layout["font"])
        normal.font.size = bullet_style.font.size = Pt(layout["body_size"])
        normal.paragraph_format.line_spacing = 1.15
        bullet_style.paragraph_format.line_spacing = 1.15
        title_style.font.size = Pt(layout["title_size"])
        title_style.paragraph_format.space_after = Pt(4)
        section_style.font.size = Pt(layout["heading_size"])
        section_style.font.color.rgb = RGBColor.from_string(layout["accent"])
        section_style.paragraph_format.space_before = Pt(layout["section_gap"])
        section_style.paragraph_format.space_after = Pt(4)
        if layout["heading_bar"]:
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), heading_background(layout))
            section_style._element.get_or_add_pPr().append(shading)

    title = draft.headline.strip() or (candidate.basics.name or "").strip() or "Resume"
    title_paragraph = document.add_paragraph(style="Title")
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if layout["title_align"] == "center" else WD_ALIGN_PARAGRAPH.LEFT
    title_paragraph.add_run(title)
    contact = _contact_line(candidate)
    if contact:
        contact_paragraph = document.add_paragraph(contact)
        contact_paragraph.alignment = title_paragraph.alignment
        contact_paragraph.paragraph_format.space_after = Pt(4)
        for run in contact_paragraph.runs:
            run.font.size = Pt(9.25)

    labels = _labels(draft, preferences.language)
    skills = _clean_items(draft.skills_section)
    education = _clean_items(draft.education_section)
    for part in layout["section_order"]:
        if part == "summary" and draft.summary.strip():
            _add_docx_heading(document, labels[part])
            document.add_paragraph(draft.summary.strip())
        elif part == "skills" and skills:
            _add_docx_heading(document, labels[part])
            document.add_paragraph(", ".join(skills))
        elif part == "experience":
            _add_docx_sections(document, labels[part], draft.experience_section)
        elif part == "projects":
            _add_docx_sections(document, labels[part], draft.project_section)
        elif part == "education" and education:
            _add_docx_heading(document, labels[part])
            for item in education:
                _add_docx_bullet(document, item)

    properties = document.core_properties
    properties.title = title
    properties.subject = "ATS friendly tailored resume"
    properties.author = "NoWorryCV"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def build_pdf(draft: ResumeDraft, candidate: CandidateProfile, preferences: Optional[ResumePreferences] = None) -> bytes:
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.pagesizes import A4, letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    preferences = preferences or ResumePreferences()
    compact = preferences.template == "ats_compact"
    layout = template_style(preferences.template)
    imported = layout["adaptation"] == "flow"
    font_name = "STSong-Light"
    try:
        pdfmetrics.getFont(font_name)
    except KeyError:
        pdfmetrics.registerFont(UnicodeCIDFont(font_name))
    if imported:
        font_name = _flow_pdf_font(layout["font"] in {"楷体", "宋体", "仿宋"})

    output = BytesIO()
    title = draft.headline.strip() or (candidate.basics.name or "").strip() or "Resume"
    document = SimpleDocTemplate(
        output,
        pagesize=A4 if paper_size(preferences) == "A4" else letter,
        rightMargin=layout["margin_x"] if imported else 0.62 * inch,
        leftMargin=layout["margin_x"] if imported else 0.62 * inch,
        topMargin=layout["margin_y"] if imported else 0.52 * inch,
        bottomMargin=layout["margin_y"] if imported else 0.52 * inch,
        title=title,
        author="NoWorryCV",
        subject="ATS friendly tailored resume",
        pageCompression=1,
        allowSplitting=1,
    )
    sample = getSampleStyleSheet()
    styles: Dict[str, ParagraphStyle] = {
        "title": ParagraphStyle(
            "ResumeTitle",
            parent=sample["Title"],
            fontName=font_name,
            fontSize=16,
            leading=18,
            textColor="#000000",
            alignment=TA_LEFT,
            spaceAfter=1 if compact else 2,
        ),
        "contact": ParagraphStyle(
            "ResumeContact",
            parent=sample["Normal"],
            fontName=font_name,
            fontSize=9.25,
            leading=11,
            textColor="#000000",
            spaceAfter=4,
        ),
        "heading": ParagraphStyle(
            "ResumeHeading",
            parent=sample["Heading2"],
            fontName=font_name,
            fontSize=11.5,
            leading=13,
            textColor="#000000",
            spaceBefore=3 if compact else 5,
            spaceAfter=1 if compact else 2,
            keepWithNext=True,
        ),
        "item": ParagraphStyle(
            "ResumeItem",
            parent=sample["Normal"],
            fontName=font_name,
            fontSize=10.25,
            leading=11.5 if compact else 12.2,
            textColor="#000000",
            spaceAfter=1,
            keepWithNext=True,
        ),
        "body": ParagraphStyle(
            "ResumeBody",
            parent=sample["Normal"],
            fontName=font_name,
            fontSize=10.25,
            leading=11.5 if compact else 12.2,
            textColor="#000000",
            spaceAfter=1 if compact else 2,
        ),
        "bullet": ParagraphStyle(
            "ResumeBullet",
            parent=sample["Normal"],
            fontName=font_name,
            fontSize=10,
            leading=11.5 if compact else 12,
            textColor="#000000",
            leftIndent=10,
            firstLineIndent=-10,
            spaceAfter=1,
        ),
    }

    if imported:
        styles["title"].fontSize = layout["title_size"]
        styles["title"].leading = layout["title_size"] * 1.2
        styles["title"].alignment = TA_CENTER if layout["title_align"] == "center" else TA_LEFT
        styles["title"].spaceAfter = 4
        styles["heading"].fontSize = layout["heading_size"]
        styles["heading"].leading = layout["heading_size"] * 1.2
        styles["heading"].textColor = "#" + layout["accent"]
        styles["heading"].spaceBefore = layout["section_gap"]
        styles["heading"].spaceAfter = 5
        if layout["heading_bar"]:
            styles["heading"].backColor = "#" + heading_background(layout)
            styles["heading"].borderPadding = 3
        for key in ("body", "bullet", "item"):
            styles[key].fontSize = layout["body_size"]
            styles[key].leading = layout["body_size"] * 1.2
            styles[key].wordWrap = "CJK"
        styles["title"].wordWrap = "CJK"

    styles["header_contact"] = ParagraphStyle("ResumeHeaderContact", parent=styles["contact"], alignment=styles["title"].alignment)

    story = [Paragraph(_escape(title), styles["title"])]
    contact = _contact_line(candidate)
    if contact:
        story.append(Paragraph(_escape(contact), styles["header_contact"]))
    labels = _labels(draft, preferences.language)
    skills = _clean_items(draft.skills_section)
    education = _clean_items(draft.education_section)
    for part in layout["section_order"]:
        if part == "summary" and draft.summary.strip():
            _add_pdf_heading(story, labels[part], styles)
            story.append(Paragraph(_escape(draft.summary.strip()), styles["body"]))
        elif part == "skills" and skills:
            _add_pdf_heading(story, labels[part], styles)
            story.append(Paragraph(_escape(", ".join(skills)), styles["body"]))
        elif part == "experience":
            _add_pdf_sections(story, labels[part], draft.experience_section, styles)
        elif part == "projects":
            _add_pdf_sections(story, labels[part], draft.project_section, styles)
        elif part == "education" and education:
            _add_pdf_heading(story, labels[part], styles)
            for item in education:
                story.append(Paragraph("- {0}".format(_escape(item)), styles["bullet"]))
    story.append(Spacer(1, 1))
    document.build(story)
    return output.getvalue()


def pdf_page_count(pdf_bytes: bytes) -> int:
    from pypdf import PdfReader

    return len(PdfReader(BytesIO(pdf_bytes)).pages)


def _labels(draft: ResumeDraft, language: str = "auto") -> Dict[str, str]:
    text = " ".join([draft.headline, draft.summary] + list(draft.skills_section))
    zh = language == "zh" or (language == "auto" and any("\u4e00" <= character <= "\u9fff" for character in text))
    if zh:
        return {
            "summary": "个人摘要",
            "skills": "技能",
            "experience": "工作经历",
            "projects": "项目经历",
            "education": "教育背景",
        }
    return {
        "summary": "Summary",
        "skills": "Skills",
        "experience": "Experience",
        "projects": "Projects",
        "education": "Education",
    }


def _contact_line(candidate: CandidateProfile) -> str:
    basics = candidate.basics
    return " | ".join(
        value.strip()
        for value in (basics.email, basics.phone, basics.location)
        if value and value.strip()
    )


def _clean_items(items: Iterable[str]) -> List[str]:
    return [item.strip() for item in items if item and item.strip()]


def _append_plain_sections(
    lines: List[str],
    label: str,
    sections: List[ResumeSectionItem],
) -> None:
    if not sections:
        return
    lines.extend(["", label])
    for section in sections:
        lines.append(section.heading.strip())
        if section.subheading and section.subheading.strip():
            lines.append(section.subheading.strip())
        lines.extend("- {0}".format(item) for item in _clean_items(section.bullets))


def _add_docx_heading(document, text: str) -> None:
    from copy import deepcopy
    from docx.oxml.ns import qn

    paragraph = document.add_paragraph(text, style="Resume Section")
    # Some Word viewers omit paragraph shading inherited from custom styles.
    properties = document.styles["Resume Section"].element.pPr
    shading = properties.find(qn("w:shd")) if properties is not None else None
    if shading is not None:
        paragraph._p.get_or_add_pPr().append(deepcopy(shading))


def _add_docx_bullet(document, text: str) -> None:
    from docx.shared import Pt

    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Pt(11)
    paragraph.paragraph_format.first_line_indent = Pt(-11)
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.add_run(text)


def _add_docx_sections(document, label: str, sections: List[ResumeSectionItem]) -> None:
    from docx.shared import Pt

    usable_sections = [section for section in sections if section.heading.strip() or _clean_items(section.bullets)]
    if not usable_sections:
        return
    _add_docx_heading(document, label)
    for section in usable_sections:
        heading = document.add_paragraph()
        heading.paragraph_format.space_before = Pt(2)
        heading.paragraph_format.space_after = Pt(0)
        heading.paragraph_format.keep_with_next = True
        run = heading.add_run(section.heading.strip())
        run.bold = True
        if section.subheading and section.subheading.strip():
            subheading = document.add_paragraph(section.subheading.strip())
            subheading.paragraph_format.space_after = Pt(1)
            subheading.paragraph_format.keep_with_next = True
            for subheading_run in subheading.runs:
                subheading_run.italic = True
                subheading_run.font.size = Pt(9.25)
        for bullet in _clean_items(section.bullets):
            _add_docx_bullet(document, bullet)


def _add_pdf_heading(story, text: str, styles: Dict[str, object]) -> None:
    from reportlab.platypus import Paragraph

    story.append(Paragraph(_escape(text), styles["heading"]))


def _add_pdf_sections(
    story,
    label: str,
    sections: List[ResumeSectionItem],
    styles: Dict[str, object],
) -> None:
    from reportlab.platypus import Paragraph

    usable_sections = [section for section in sections if section.heading.strip() or _clean_items(section.bullets)]
    if not usable_sections:
        return
    _add_pdf_heading(story, label, styles)
    for section in usable_sections:
        if section.heading.strip():
            story.append(Paragraph(_escape(section.heading.strip()), styles["item"]))
        if section.subheading and section.subheading.strip():
            story.append(Paragraph(_escape(section.subheading.strip()), styles["contact"]))
        for bullet in _clean_items(section.bullets):
            story.append(Paragraph("- {0}".format(_escape(bullet)), styles["bullet"]))


def _escape(value: str) -> str:
    return html.escape(value or "", quote=False).replace("\n", "<br/>")
