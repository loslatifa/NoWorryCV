from io import BytesIO
import json
from pathlib import Path
import subprocess
import sys
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from fastapi.testclient import TestClient
from pypdf import PdfReader
import pytest

from backend.app.main import create_app
from backend.app.schemas.preferences import ResumePreferences
from backend.app.schemas.strategy import ResumeSectionItem
from backend.app.services.resume_export import build_docx, build_pdf
from backend.app.services.runs.store import get_run_store
from backend.app.services.template_catalog import TEMPLATES
from backend.tests.test_candidate_profiles import save_archive, wait_for_run
from backend.tests.test_draft_editing import _completed_run
from backend.tests.test_resume_export import _sample_resume

SUPPLIED = [key for key, item in TEMPLATES.items() if item["adaptation"] == "flow" and item["active"]]
ROOT = Path(__file__).resolve().parents[2]


@pytest.mark.parametrize("template_id", list(TEMPLATES))
def test_every_template_exports_current_content_without_sample_content(template_id):
    candidate, draft = _sample_resume()
    draft.project_section = [ResumeSectionItem(heading="当前用户项目", bullets=["使用 Power BI 构建看板，配合 CI/CD 交付。"])]
    preferences = ResumePreferences(template=template_id, region="CN", language="zh")
    docx = build_docx(draft, candidate, preferences)
    pdf = build_pdf(draft, candidate, preferences)
    document = Document(BytesIO(docx))
    word_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    reader = PdfReader(BytesIO(pdf))
    pdf_text = "\n".join(page.extract_text() or "" for page in reader.pages)
    for text in ("张三 产品经理", "使用 SQL 分析增长漏斗", "当前用户项目", "Power BI", "CI/CD", "示例大学"):
        assert text in word_text
        assert text in pdf_text
    assert "resume@example.com" not in word_text + pdf_text
    assert "示例姓名" not in word_text + pdf_text
    assert float(reader.pages[0].mediabox.width) == pytest.approx(595.28, abs=.1)
    with ZipFile(BytesIO(docx)) as archive:
        assert not any(name.startswith("word/media/") for name in archive.namelist())
    if template_id in SUPPLIED:
        style = TEMPLATES[template_id]
        assert document.styles["Normal"].font.name == style["font"]
        assert document.styles["Title"].font.size.pt == style["title_size"]
        heading = document.styles["Resume Section"]
        if style["heading_bar"]:
            assert heading.element.find(".//" + qn("w:shd")) is not None
            assert str(heading.font.color.rgb) == style["accent"]
            assert str(heading.font.color.rgb) != "FFFFFF"


def test_catalog_previews_and_thumbnails_are_available_without_source_directory():
    client = TestClient(create_app())
    payload = client.get("/api/v1/resume-templates").json()
    assert SUPPLIED
    assert {item["id"] for item in payload["templates"]} == set(TEMPLATES)
    assert "/Users/" not in json.dumps(payload)
    for template_id in SUPPLIED:
        response = client.get(f"/api/v1/resume-templates/{template_id}/preview")
        assert response.status_code == 200
        assert response.headers["x-resume-template"] == template_id
        assert "示例姓名" in PdfReader(BytesIO(response.content)).pages[0].extract_text()
        image = client.get(f"/app/templates/{template_id}.png")
        assert image.status_code == 200
        assert image.content.startswith(b"\x89PNG")
    assert client.get("/api/v1/resume-templates/not-a-template/preview").status_code == 422


def test_export_template_override_keeps_saved_run_and_draft_unchanged():
    job = _completed_run()
    client = TestClient(create_app())
    request = {"draft": job.result.final_package.draft.model_dump(), "template": SUPPLIED[0]}
    path = f"/api/v1/tailor-runs/{job.run_id}"
    headers = {"X-Run-Token": job.access_token}
    before = get_run_store().get(job.run_id).result.model_dump()
    assert client.post(path + "/export-preview", json=request).status_code == 404
    preview = client.post(path + "/export-preview", json=request, headers=headers)
    assert preview.status_code == 200
    assert preview.headers["x-resume-template"] == SUPPLIED[0]
    for format_name in ("docx", "pdf"):
        response = client.post(path + f"/export/{format_name}", json=request, headers=headers)
        assert response.status_code == 200
        assert response.headers["x-resume-template"] == SUPPLIED[0]
    after = get_run_store().get(job.run_id).result.model_dump()
    assert before == after
    assert client.post(path + "/export-preview", json={**request, "template": "unknown"}, headers=headers).status_code == 422
    fallback = client.post(path + "/export-preview", json={"draft": request["draft"]}, headers=headers)
    assert fallback.headers["x-resume-template"] == before["preferences"]["template"]


def test_supplied_template_preference_inherited_and_previous_run_stays_frozen():
    with TestClient(create_app()) as client:
        archive, headers = save_archive(client)
        endpoint = f'/api/v1/candidate-profiles/{archive["id"]}'
        preferences = {"template": SUPPLIED[0], "region": "CN"}
        response = client.patch(endpoint + "/preferences", headers=headers,
                                json={"preferences": preferences, "expected_preferences_revision": 1})
        assert response.status_code == 200
        job = client.post(endpoint + "/jobs", headers=headers,
                          json={"jd_text": "数据分析师 SQL", "expected_revision": 1, "max_iterations": 1}).json()
        result = wait_for_run(client, job)["result"]
        assert result["preferences"]["template"] == SUPPLIED[0]
        client.patch(endpoint + "/preferences", headers=headers,
                     json={"preferences": {"template": SUPPLIED[-1]}, "expected_preferences_revision": 2})
        stored = get_run_store().get(job["run_id"]).result
        assert stored.preferences.template == SUPPLIED[0]


@pytest.mark.parametrize("template_id", [SUPPLIED[0], SUPPLIED[-1]])
def test_long_bilingual_content_flows_across_pages_without_truncation(template_id):
    candidate, draft = _sample_resume()
    draft.experience_section[0].bullets = [
        f"唯一条目 {index:03d} 使用 SQL 与 Python 完成 customer analytics 和跨团队交付。" * 2
        for index in range(90)
    ]
    pdf = build_pdf(draft, candidate, ResumePreferences(template=template_id, region="US"))
    reader = PdfReader(BytesIO(pdf))
    assert len(reader.pages) > 2
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    for index in range(90):
        assert f"唯一条目 {index:03d}" in text
    assert float(reader.pages[0].mediabox.width) == 612


def test_import_omits_source_text_and_retains_old_versions(tmp_path):
    directory = tmp_path / "templates"
    directory.mkdir()
    source = directory / "01号-测试模板.docx"
    document = Document()
    document.add_paragraph("PRIVATE_EXAMPLE_SHOULD_NOT_BECOME_CANDIDATE_CONTENT")
    document.add_paragraph("教育背景")
    document.save(source)
    (directory / "unsupported.pdf").write_bytes(b"not a DOCX")
    output = tmp_path / "catalog.json"
    command = [sys.executable, str(ROOT / "scripts/import_resume_templates.py"), str(directory), "--output", str(output)]
    subprocess.run(command, check=True, capture_output=True)
    first = json.loads(output.read_text())
    assert len(first["templates"]) == 1
    assert first["issues"][0]["filename"] == "unsupported.pdf"
    assert "PRIVATE_EXAMPLE" not in output.read_text()
    original_id = first["templates"][0]["id"]
    document.add_paragraph("New template revision")
    document.save(source)
    subprocess.run(command, check=True, capture_output=True)
    second = json.loads(output.read_text())
    assert len(second["templates"]) == 2
    assert next(item for item in second["templates"] if item["id"] == original_id)["active"] is False
